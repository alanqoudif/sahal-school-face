from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.seats import seat_label


def _safe_sheet_name(title: str) -> str:
    cleaned = "".join(ch for ch in title if ch not in r"[]:*?/\\")
    return (cleaned or "صف")[:31]


def attendance_workbook(day, classrooms, present_by_student: dict) -> BytesIO:
    book = Workbook()
    book.remove(book.active)
    header_fill = PatternFill("solid", fgColor="059669")
    header_font = Font(color="FFFFFF", bold=True)
    present_fill = PatternFill("solid", fgColor="D1FAE5")
    absent_fill = PatternFill("solid", fgColor="FEF3C7")

    if not classrooms:
        sheet = book.create_sheet("لا توجد صفوف")
        sheet.append(["ما فيه صفوف بعد"])
    else:
        used_names: set[str] = set()
        for classroom in classrooms:
            name = _safe_sheet_name(f"{classroom.grade} - {classroom.section}")
            original = name
            index = 2
            while name in used_names:
                name = f"{original[:28]}-{index}"
                index += 1
            used_names.add(name)
            sheet = book.create_sheet(name)
            sheet.append(
                [
                    "الاسم",
                    "رقم الطالب",
                    "الصف",
                    "الشعبة",
                    "المقعد",
                    "الحالة",
                    "وقت الدخول",
                    "وضوح الصورة",
                    "مصدر التسجيل",
                ]
            )
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font

            students = sorted(classroom.students, key=lambda item: (item.seat_code or "zzz", item.name))
            for student in students:
                record = present_by_student.get(student.id)
                status = "حاضر" if record else "غائب"
                time_text = record.checked_in_at.strftime("%H:%M") if record else ""
                sheet.append(
                    [
                        student.name,
                        student.student_number,
                        classroom.grade,
                        classroom.section,
                        seat_label(student.seat_code),
                        status,
                        time_text,
                        getattr(record, "quality", "") if record else "",
                        "يدوي" if record and getattr(record, "source", "") == "manual" else ("كاميرا" if record else ""),
                    ]
                )
                fill = present_fill if record else absent_fill
                for cell in sheet[sheet.max_row]:
                    cell.fill = fill
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 18

    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output
